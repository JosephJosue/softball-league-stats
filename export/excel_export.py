"""Excel export helpers.

Builds styled .xlsx workbooks in memory (bytes) for Streamlit download buttons.
UI-agnostic: callers pass {sheet_name: DataFrame}; this module handles writing,
header styling, column widths, and a frozen header row.
"""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="1F7A3D")  # grass green
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CENTER = Alignment(horizontal="center", vertical="center")

# UUID/bookkeeping columns we don't want in an export.
_DROP_COLS = ["id", "player_id", "team_id", "game_id", "position_id", "created_at"]


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    """Drop id/bookkeeping columns for a tidy export."""
    if df is None or df.empty:
        return pd.DataFrame() if df is None else df
    return df.drop(columns=[c for c in _DROP_COLS if c in df.columns])


def _style_sheet(worksheet, df: pd.DataFrame) -> None:
    for col_idx, col in enumerate(df.columns, start=1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        sample = df[col].astype(str).head(200)
        width = max([len(str(col)), *(len(v) for v in sample)]) if len(sample) else len(str(col))
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 42)
    worksheet.freeze_panes = "A2"


def build_workbook(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Return xlsx bytes; one sheet per {name: DataFrame} (id columns dropped)."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        wrote_any = False
        for name, df in sheets.items():
            clean = _clean(df)
            if clean is None or clean.empty:
                continue
            sheet_name = name[:31]  # Excel's 31-char sheet-name limit
            clean.to_excel(writer, sheet_name=sheet_name, index=False)
            _style_sheet(writer.sheets[sheet_name], clean)
            wrote_any = True
        if not wrote_any:
            pd.DataFrame({"info": ["No data"]}).to_excel(writer, sheet_name="Empty", index=False)
    return buffer.getvalue()
