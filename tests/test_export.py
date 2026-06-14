"""Tests for the Excel export helper."""

import io

import pandas as pd
from openpyxl import load_workbook

from export import excel_export


def test_workbook_is_valid_xlsx_and_drops_ids():
    df = pd.DataFrame([{"id": "x", "team_id": "t", "player_name": "A", "ab": 4, "h": 2}])
    data = excel_export.build_workbook({"Players": df})
    assert data[:2] == b"PK"  # xlsx is a zip
    wb = load_workbook(io.BytesIO(data))
    headers = [c.value for c in wb["Players"][1]]
    assert "player_name" in headers
    assert "id" not in headers and "team_id" not in headers
    assert wb["Players"].freeze_panes == "A2"


def test_empty_sheets_produce_placeholder():
    data = excel_export.build_workbook({"Players": pd.DataFrame()})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Empty"]


def test_long_sheet_name_truncated():
    df = pd.DataFrame([{"x": 1}])
    name = "A really long sheet name beyond limit"
    data = excel_export.build_workbook({name: df})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == name[:31]
